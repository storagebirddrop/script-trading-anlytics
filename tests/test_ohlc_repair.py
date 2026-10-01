#!/usr/bin/env python3
"""
Tests for the High/Low/Volume repair path and for ADX/Bollinger being derived
from history instead of read from (non-existent) history columns.
"""

import numpy as np
import openpyxl
import pandas as pd

import calculate_metrics as cm
import crypto_tracker
import update_history as uh
from trading_utils.excel_utils import ensure_excel_headers


def _history(n=60, with_ohlc=True, last_missing=False):
    close = 100 + np.cumsum(np.sin(np.arange(n) / 3.0))
    df = pd.DataFrame({
        'Date': pd.date_range('2026-01-01', periods=n).strftime('%Y-%m-%d'),
        'Asset': 'BTC',
        'Timeframe': '1d',
        'Price': close,
    })
    if with_ohlc:
        df['High'] = close + 1.0
        df['Low'] = close - 1.0
        df['Volume'] = 1000.0
        if last_missing:
            df.loc[df.index[-1], ['High', 'Low', 'Volume']] = np.nan
    return df


class TestAdxBbFromHistory:
    def test_all_values_present_with_enough_history(self):
        out = cm._adx_bb_from_history(_history())
        assert out['adx'] is not None and 0 <= out['adx'] <= 100
        assert out['bb_pct_b'] is not None
        assert out['bb_bandwidth'] is not None and out['bb_bandwidth'] > 0

    def test_adx_none_when_latest_bar_has_no_high_low(self):
        out = cm._adx_bb_from_history(_history(last_missing=True))
        assert out['adx'] is None
        # Bollinger only needs Price, so it must survive an OHLC gap
        assert out['bb_pct_b'] is not None

    def test_adx_none_without_ohlc_columns(self):
        out = cm._adx_bb_from_history(_history(with_ohlc=False))
        assert out['adx'] is None
        assert out['bb_pct_b'] is not None

    def test_all_none_with_too_little_history(self):
        out = cm._adx_bb_from_history(_history(n=10))
        assert out == {'adx': None, 'bb_pct_b': None, 'bb_bandwidth': None}

    def test_current_metrics_expose_adx_and_bb(self):
        df = _history()
        df['EMA21'] = df['Price'] - 0.5
        df['ATR'] = 1.4
        df['RSI'] = 50.0
        df['RSI_Z_Score'] = 0.0
        df['ATR_Distance'] = 0.35
        df['Pct_Above_EMA'] = 0.5
        current = cm.calculate_current_metrics(df)['BTC']['1d']['current']
        assert current['adx'] is not None
        assert current['bb_pct_b'] is not None
        assert current['bb_bandwidth'] is not None


class TestLegacyUnnamedColumns:
    def test_name_unnamed_columns_by_position(self):
        df = pd.DataFrame(columns=[
            'Date', 'Asset', 'Timeframe', 'Price', 'EMA21', 'ATR', 'RSI',
            'RSI_Z_Score', 'ATR_Distance', 'Pct_Above_EMA',
            'Unnamed: 10', 'Unnamed: 11', 'Unnamed: 12',
        ])
        out = uh.name_unnamed_columns(df)
        assert list(out.columns[-3:]) == ['High', 'Low', 'Volume']

    def test_merge_fills_gaps_and_keeps_existing_values(self):
        df = pd.DataFrame({
            'High': [5.0, np.nan], 'Low': [4.0, np.nan], 'Volume': [10.0, np.nan],
            'Unnamed: 10': [9.0, 7.0], 'Unnamed: 11': [8.0, 6.0], 'Unnamed: 12': [99.0, 20.0],
        })
        out = uh.merge_legacy_unnamed_columns(df)
        assert out['High'].tolist() == [5.0, 7.0]
        assert out['Low'].tolist() == [4.0, 6.0]
        assert out['Volume'].tolist() == [10.0, 20.0]
        assert not any(str(c).startswith('Unnamed') for c in out.columns)

    def test_merge_creates_columns_when_absent(self):
        df = pd.DataFrame({'Unnamed: 10': [1.0], 'Unnamed: 11': [0.5], 'Unnamed: 12': [3.0]})
        out = uh.merge_legacy_unnamed_columns(df)
        assert out.iloc[0].to_dict() == {'High': 1.0, 'Low': 0.5, 'Volume': 3.0}


class TestEnsureExcelHeaders:
    def test_adds_only_missing_headers(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        for col, h in enumerate(['Date', 'Asset', 'Timeframe'], 1):
            ws.cell(row=1, column=col, value=h)
        added = ensure_excel_headers(ws, ['Date', 'Asset', 'Timeframe', 'Price', 'High'])
        assert added == 2
        assert [c.value for c in ws[1]][:5] == ['Date', 'Asset', 'Timeframe', 'Price', 'High']

    def test_complete_header_row_is_untouched(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        headers = ['Date', 'Asset']
        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h)
        assert ensure_excel_headers(ws, headers) == 0


class TestFindMissingOhlc:
    def test_flags_records_without_high_or_low(self):
        records = [
            {'Asset': 'BTC', 'Timeframe': '1d', 'High': 2.0, 'Low': 1.0},
            {'Asset': 'ETH', 'Timeframe': '1w', 'High': None, 'Low': 1.0},
            {'Asset': 'XLM', 'Timeframe': '1M', 'High': 2.0, 'Low': None},
        ]
        assert crypto_tracker.find_missing_ohlc(records) == ['ETH (1w)', 'XLM (1M)']

    def test_clean_records_report_nothing(self):
        records = [{'Asset': 'BTC', 'Timeframe': '1d', 'High': 2.0, 'Low': 1.0}]
        assert crypto_tracker.find_missing_ohlc(records) == []


class TestRefreshOpenBars:
    def _frame(self, price, high=2.0):
        return pd.DataFrame({
            'Date': ['2026-09-28', '2026-09-21'],
            'Asset': ['XLM', 'XLM'],
            'Timeframe': ['1w', '1w'],
            'Price': [price, 1.0],
            'EMA21': [0.9, 0.9], 'ATR': [0.1, 0.1], 'RSI': [50.0, 50.0],
            'RSI_Z_Score': [0.0, 0.0], 'ATR_Distance': [0.5, 0.5], 'Pct_Above_EMA': [1.0, 1.0],
            'High': [high, 1.1], 'Low': [1.0, 0.9], 'Volume': [10.0, 10.0],
        })

    def test_latest_bar_gets_newer_values(self):
        history = self._frame(price=1.50)
        excel = self._frame(price=1.60, high=2.5)
        out, n = uh.refresh_open_bars(history, excel)
        row = out[out['Date'] == '2026-09-28'].iloc[0]
        assert n == 1
        assert row['Price'] == 1.60 and row['High'] == 2.5

    def test_older_closed_bars_are_left_alone(self):
        history = self._frame(price=1.50)
        excel = self._frame(price=1.60)
        excel.loc[excel['Date'] == '2026-09-21', 'Price'] = 99.0
        out, _ = uh.refresh_open_bars(history, excel)
        assert out[out['Date'] == '2026-09-21'].iloc[0]['Price'] == 1.0

    def test_unfinished_bar_without_price_does_not_overwrite(self):
        history = self._frame(price=1.50)
        excel = self._frame(price=np.nan, high=9.0)
        out, n = uh.refresh_open_bars(history, excel)
        assert n == 0
        assert out[out['Date'] == '2026-09-28'].iloc[0]['High'] == 2.0

    def test_identical_values_report_no_update(self):
        history = self._frame(price=1.50)
        out, n = uh.refresh_open_bars(history, self._frame(price=1.50))
        assert n == 0

    def test_nan_never_overwrites_stored_value(self):
        history = self._frame(price=1.50)
        excel = self._frame(price=1.60)
        excel.loc[excel['Date'] == '2026-09-28', 'High'] = np.nan
        out, _ = uh.refresh_open_bars(history, excel)
        assert out[out['Date'] == '2026-09-28'].iloc[0]['High'] == 2.0


class TestWriteToExcelUpsert:
    def _record(self, price, high=2.0):
        return {
            'Date': '2026-09-28', 'Asset': 'XLM', 'Timeframe': '1w', 'Price': price,
            'EMA21': 0.9, 'ATR': 0.1, 'RSI': 50.0, 'RSI_Z_Score': 0.0,
            'ATR_Distance': 0.5, 'Pct_Above_EMA': 1.0, 'High': high, 'Low': 1.0, 'Volume': 10.0,
        }

    def _run(self, tmp_path, monkeypatch, *batches):
        path = tmp_path / 'tracker.xlsx'
        monkeypatch.setattr(crypto_tracker, '_SPREADSHEET', path)
        for batch in batches:
            crypto_tracker.write_to_excel(batch)
        ws = openpyxl.load_workbook(path)['Data']
        return [[c.value for c in row] for row in ws.iter_rows(min_row=2)]

    def test_second_write_refreshes_the_open_bar_instead_of_skipping(self, tmp_path, monkeypatch):
        rows = self._run(tmp_path, monkeypatch, [self._record(1.50)], [self._record(1.60, high=2.5)])
        assert len(rows) == 1
        assert rows[0][3] == 1.60 and rows[0][10] == 2.5

    def test_unfinished_bar_does_not_overwrite(self, tmp_path, monkeypatch):
        rows = self._run(tmp_path, monkeypatch, [self._record(1.50)], [self._record(None, high=9.0)])
        assert rows[0][3] == 1.50 and rows[0][10] == 2.0

    def test_headers_include_high_low_volume(self, tmp_path, monkeypatch):
        path = tmp_path / 'tracker.xlsx'
        monkeypatch.setattr(crypto_tracker, '_SPREADSHEET', path)
        crypto_tracker.write_to_excel([self._record(1.50)])
        ws = openpyxl.load_workbook(path)['Data']
        assert [c.value for c in ws[1]][-3:] == ['High', 'Low', 'Volume']

    def test_legacy_workbook_without_headers_is_repaired(self, tmp_path, monkeypatch):
        path = tmp_path / 'tracker.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Data'
        for col, h in enumerate(crypto_tracker._EXCEL_HEADERS[:10], 1):
            ws.cell(row=1, column=col, value=h)
        wb.save(path)
        monkeypatch.setattr(crypto_tracker, '_SPREADSHEET', path)
        crypto_tracker.write_to_excel([self._record(1.50)])
        ws = openpyxl.load_workbook(path)['Data']
        assert [c.value for c in ws[1]][-3:] == ['High', 'Low', 'Volume']


class TestGetDataSkipsUnfinishedBar:
    def _ohlcv(self, n=60, last_close_nan=True):
        idx = pd.date_range('2026-07-01', periods=n)
        close = pd.Series(100 + np.cumsum(np.sin(np.arange(n) / 3.0)), index=idx)
        df = pd.DataFrame({'open': close, 'high': close + 1, 'low': close - 1,
                           'close': close, 'volume': 1000.0}, index=idx)
        if last_close_nan:
            df.iloc[-1, df.columns.get_loc('close')] = np.nan
        return df

    def test_trailing_nan_close_is_dropped(self, monkeypatch):
        df = self._ohlcv()
        monkeypatch.setattr(crypto_tracker, 'fetch_ohlcv_yahoo', lambda *a, **k: df)
        rec = crypto_tracker.get_data('BTC', '1d')
        assert rec is not None
        assert rec['Date'] == df.index[-2].strftime('%Y-%m-%d')
        assert rec['Price'] == df['close'].iloc[-2]
        assert rec['High'] is not None and rec['Low'] is not None

    def test_complete_latest_bar_is_used_as_is(self, monkeypatch):
        df = self._ohlcv(last_close_nan=False)
        monkeypatch.setattr(crypto_tracker, 'fetch_ohlcv_yahoo', lambda *a, **k: df)
        rec = crypto_tracker.get_data('BTC', '1d')
        assert rec['Date'] == df.index[-1].strftime('%Y-%m-%d')

    def test_all_nan_close_returns_none(self, monkeypatch):
        df = self._ohlcv()
        df['close'] = np.nan
        monkeypatch.setattr(crypto_tracker, 'fetch_ohlcv_yahoo', lambda *a, **k: df)
        assert crypto_tracker.get_data('BTC', '1d') is None
